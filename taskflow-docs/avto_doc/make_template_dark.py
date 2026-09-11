"""Создаёт второй .docx-шаблон — тёмный 'Midnight', резко контрастный первому."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


PALETTE = {
    "bg":           "1A1613",   # почти чёрный, тёплый
    "primary":      "C99A6D",   # светлое золото
    "primary_dark": "D9B587",   # для крупных заголовков
    "accent":       "E0B87A",   # тёплый мёд
    "text":         "E8E2D6",   # крем
    "text_muted":   "A89E8E",   # серо-бежевый
    "code_bg":      "2A2620",
    "code_fg":      "E0B87A",
    "border":       "3A322A",
}


def rgb(h):
    return RGBColor.from_string(h)


def set_font(style, name, size, bold=False, italic=False, color=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    if color:
        f.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(attr), name)


def add_border(paragraph, color, position='bottom', size=6, space=6):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    el = OxmlElement(f'w:{position}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(size))
    el.set(qn('w:space'), str(space))
    el.set(qn('w:color'), color)
    pbdr.append(el)
    pPr.append(pbdr)


def add_field(paragraph, code):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve')
    i.text = code
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for el in (b, i, s, e):
        run._r.append(el)
    return run


def set_page_background(doc, fill_hex):
    """Красит фон страницы во всём документе."""
    doc_el = doc.element
    bg = doc_el.find(qn('w:background'))
    if bg is None:
        bg = OxmlElement('w:background')
        doc_el.insert(0, bg)
    bg.set(qn('w:color'), fill_hex)

    settings = doc.settings.element
    if settings.find(qn('w:displayBackgroundShape')) is None:
        settings.append(OxmlElement('w:displayBackgroundShape'))


# ---------- Стили ----------
def setup_styles(doc):
    styles = doc.styles

    normal = styles['Normal']
    set_font(normal, 'Calibri', 11, color=PALETTE['text'])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    specs = [
        ('Heading 1', 24, 'primary_dark', 28, 10, 'Georgia'),
        ('Heading 2', 17, 'primary',      20, 8,  'Georgia'),
        ('Heading 3', 13, 'accent',       14, 6,  'Georgia'),
        ('Heading 4', 11, 'text',         10, 4,  'Georgia'),
    ]
    for name, size, key, before, after, font in specs:
        s = styles[name]
        set_font(s, font, size, bold=True, color=PALETTE[key])
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    try:
        code = styles.add_style('Code', 1)
    except Exception:
        code = styles['Code']
    set_font(code, 'Consolas', 10, color=PALETTE['code_fg'])
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.4)

    q = styles['Quote']
    set_font(q, 'Georgia', 12, italic=True, color=PALETTE['primary'])
    q.paragraph_format.left_indent = Cm(0.6)
    q.paragraph_format.right_indent = Cm(0.6)


# ---------- Титульный лист ----------
def add_title_page(doc):
    # Верхняя золотая линия
    top = doc.add_paragraph()
    top.paragraph_format.space_before = Pt(60)
    add_border(top, PALETTE['primary'], 'bottom', 12, 0)
    r = top.add_run(' ')
    r.font.size = Pt(1)

    # Кикер сверху
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kr = kicker.add_run('TASKFLOW  ·  DOCUMENTATION')
    kr.font.name = 'Calibri'
    kr.font.size = Pt(10)
    kr.font.bold = True
    kr.font.color.rgb = rgb(PALETTE['primary'])
    kicker.paragraph_format.space_before = Pt(24)

    # Большой заголовок
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(100)
    r = p.add_run('{{ title }}')
    r.font.name = 'Georgia'
    r.font.size = Pt(48)
    r.font.bold = True
    r.font.color.rgb = rgb(PALETTE['primary_dark'])

    # Подзаголовок
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run('{{ subtitle }}')
    r2.font.name = 'Georgia'
    r2.font.size = Pt(18)
    r2.font.italic = True
    r2.font.color.rgb = rgb(PALETTE['primary'])

    # Золотая линия под подзаголовком
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(16)
    add_border(sep, PALETTE['accent'], 'bottom', 20, 0)
    r = sep.add_run(' ')
    r.font.size = Pt(1)

    # Низ титульного листа
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(280)
    r3 = p3.add_run('ВЕРСИЯ {{ version }}')
    r3.font.name = 'Calibri'
    r3.font.size = Pt(11)
    r3.font.bold = True
    r3.font.color.rgb = rgb(PALETTE['accent'])

    p4 = doc.add_paragraph()
    r4 = p4.add_run('{{ author }}\n{{ date }}')
    r4.font.size = Pt(11)
    r4.font.color.rgb = rgb(PALETTE['text_muted'])

    doc.add_page_break()


# ---------- Оглавление ----------
def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run('Содержание')
    r.font.name = 'Georgia'
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = rgb(PALETTE['primary_dark'])
    p.paragraph_format.space_after = Pt(12)
    add_border(p, PALETTE['accent'], 'bottom', 6, 8)

    toc_p = doc.add_paragraph()
    add_field(toc_p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()


# ---------- Колонтитулы ----------
def add_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = hp.add_run('—  {{ title }}  ·  v{{ version }}  —')
    r1.font.name = 'Georgia'
    r1.font.size = Pt(9)
    r1.font.italic = True
    r1.font.color.rgb = rgb(PALETTE['primary'])
    add_border(hp, PALETTE['border'], 'bottom', 4, 4)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('·  ')
    r.font.color.rgb = rgb(PALETTE['primary'])
    r.font.size = Pt(9)
    add_field(fp, 'PAGE')
    fp.add_run('  ·  ')
    add_field(fp, 'NUMPAGES')
    fp.add_run('  ·')
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = rgb(PALETTE['primary'])


# ---------- Основной контент ----------
def add_content_block(doc):
    doc.add_paragraph('{{ section.title }}', style='Heading 1')
    doc.add_paragraph('{{ section.body }}')
    doc.add_paragraph('')


# ---------- Сборка ----------
def make_template(output_path='templates/template-dark.docx'):
    doc = Document()
    set_page_background(doc, PALETTE['bg'])
    setup_styles(doc)
    add_title_page(doc)
    add_toc(doc)

    doc.add_paragraph('{% for section in sections %}', style='Normal')
    add_content_block(doc)
    doc.add_paragraph('{% endfor %}', style='Normal')

    for section in doc.sections:
        add_header_footer(section)

    doc.save(output_path)
    print(f'Тёмный шаблон создан: {output_path}')


if __name__ == '__main__':
    import os
    os.makedirs('templates', exist_ok=True)
    make_template()