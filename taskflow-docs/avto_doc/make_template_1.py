"""Создаёт .docx-шаблон с оформлением TaskFlow для использования с docxtpl."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


PALETTE = {
    "primary":      "9C5B45",   # терракота
    "primary_dark": "6B4030",   # тёмный коричневый
    "accent":       "B8894D",   # мёд
    "text":         "2E2B27",   # графит
    "text_muted":   "5A544C",   # серо-коричневый
    "bg_soft":      "F1ECE3",   # крем для кода
    "code_fg":      "6B3A28",   # код
    "border":       "E5DDD0",   # границы
}


def rgb(h):
    return RGBColor.from_string(h)


def set_font(style, name, size, bold=False, italic=False, color=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    if color:
        f.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(attr), name)


def add_border(paragraph, color, position='bottom', size=6, space=6):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    el = OxmlElement(f'w:{position}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(size))
    el.set(qn('w:space'), str(space))
    el.set(qn('w:color'), color)
    pbdr.append(el)
    pPr.append(pbdr)


def shade(paragraph, fill_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)


def add_field(paragraph, code):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve')
    i.text = code
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for el in (b, i, s, e):
        run._r.append(el)
    return run


# ---------- Стили ----------
def setup_styles(doc):
    styles = doc.styles

    normal = styles['Normal']
    set_font(normal, 'Calibri', 11, color=PALETTE['text'])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    specs = [
        ('Heading 1', 20, 'primary_dark', 24, 8),
        ('Heading 2', 15, 'primary',      18, 6),
        ('Heading 3', 12, 'text',         14, 4),
        ('Heading 4', 11, 'text_muted',   10, 2),
    ]
    for name, size, key, before, after in specs:
        s = styles[name]
        set_font(s, 'Calibri Light', size, bold=True, color=PALETTE[key])
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    # Code
    try:
        code = styles.add_style('Code', 1)  # 1 = paragraph
    except Exception:
        code = styles['Code']
    set_font(code, 'Consolas', 10, color=PALETTE['code_fg'])
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.4)

    # Quote
    q = styles['Quote']
    set_font(q, 'Calibri', 11, italic=True, color=PALETTE['text_muted'])
    q.paragraph_format.left_indent = Cm(0.6)
    q.paragraph_format.right_indent = Cm(0.6)


# ---------- Титульный лист ----------
def add_title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run('{{ title }}')
    r.font.name = 'Calibri Light'
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = rgb(PALETTE['primary_dark'])

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run('{{ subtitle }}')
    r2.font.name = 'Calibri Light'
    r2.font.size = Pt(16)
    r2.font.color.rgb = rgb(PALETTE['primary'])

    sep = doc.add_paragraph()
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_border(sep, PALETTE['accent'], size=8, space=2)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(260)
    r3 = p3.add_run('Версия {{ version }}\n{{ author }}\n{{ date }}')
    r3.font.size = Pt(11)
    r3.font.color.rgb = rgb(PALETTE['text_muted'])

    doc.add_page_break()


# ---------- Оглавление ----------
def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run('Содержание')
    r.font.name = 'Calibri Light'
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = rgb(PALETTE['primary_dark'])
    p.paragraph_format.space_after = Pt(12)

    toc_p = doc.add_paragraph()
    add_field(toc_p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()


# ---------- Колонтитулы ----------
def add_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.text = ''
    r1 = hp.add_run('{{ title }}')
    r1.font.size = Pt(9)
    r1.font.color.rgb = rgb(PALETTE['text_muted'])
    hp.add_run('\t\t')
    r2 = hp.add_run('v{{ version }}')
    r2.font.size = Pt(9)
    r2.font.color.rgb = rgb(PALETTE['text_muted'])
    add_border(hp, PALETTE['border'], 'bottom', 4)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, 'PAGE')
    fp.add_run(' из ')
    add_field(fp, 'NUMPAGES')
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = rgb(PALETTE['text_muted'])


# ---------- Основной контент ----------
def add_content_block(doc):
    doc.add_paragraph('{{ section.title }}', style='Heading 1')
    doc.add_paragraph('{{ section.body }}')
    doc.add_paragraph('')


# ---------- Сборка ----------
def make_template(output_path='templates/template.docx'):
    doc = Document()
    setup_styles(doc)
    add_title_page(doc)
    add_toc(doc)

    # Цикл docxtpl по секциям
    doc.add_paragraph('{% for section in sections %}', style='Normal')
    add_content_block(doc)
    doc.add_paragraph('{% endfor %}', style='Normal')

    for section in doc.sections:
        add_header_footer(section)

    doc.save(output_path)
    print(f'Шаблон создан: {output_path}')


if __name__ == '__main__':
    import os
    os.makedirs('templates', exist_ok=True)
    make_template()