"""Базовый документ ГОСТ-генератора — без primer.docx.

Строится программно из пустого шаблона python-docx и содержит только
структуру, которую ждёт ``avto_doc/generate_single.py``:

* стили: заголовки ``1..9`` (по уровням), ``11/21/31`` (содержание),
  ``ae`` (заголовок содержания), ``a8/aa`` (колонтитулы), ``a0``/``a2``
  (перечни), ``af3/af4`` (подписи), ``af5`` (таблицы), ``af0`` (ссылки);
* нумерация: абзац ``Normal``/docDefaults по ГОСТ (12 пт, интервал 1,5,
  по ширине), маркированные пункты — «•»;
* элементы документа: заголовок «Изменения в документе», скелет листа
  регистрации изменений (2 строки), скелет поля содержания (sdt).

Всё оформление (титул, поля, колонтитулы, порядок частей) применяется
потом в ``titlepage.py`` и ``gost_styles.py``.
"""

from __future__ import annotations

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

_REV_HEADING = "Изменения в документе"
_REV_HEADERS = ("Версия", "Дата", "Автор", "Описание")
_TOC_TAB = '<w:tab w:val="right" w:leader="dot" w:pos="9911"/>'
_WARNING_XML: str | None = None


def _has_style(doc, style_id: str) -> bool:
    return any(
        s.element.get(qn("w:styleId")) == style_id for s in doc.styles
    )


def _add_style(doc, xml: str):
    doc.styles.element.append(parse_xml(xml))


def _install_styles(doc) -> None:
    rename = {
        **{f"Heading{i}": str(i) for i in range(1, 10)},
        "Header": "a8",
        "Footer": "aa",
        "TOCHeading": "ae",
    }
    for s in doc.styles:
        new_id = rename.get(s.element.get(qn("w:styleId")))
        if new_id is not None:
            s.element.set(qn("w:styleId"), new_id)

    para = '<w:style {ns} w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/><w:basedOn w:val="Normal"/><w:qFormat/>{pPr}{rPr}</w:style>'
    char = '<w:style {ns} w:type="character" w:styleId="af0"><w:name w:val="Hyperlink"/><w:qFormat/><w:rPr><w:color w:val="0563C3"/><w:u w:val="single"/></w:rPr></w:style>'

    specs = (
        ("11", "toc 1", "", "<w:rPr><w:sz w:val=\"24\"/><w:szCs w:val=\"24\"/></w:rPr>"),
        ("21", "toc 2", "", "<w:rPr><w:sz w:val=\"24\"/><w:szCs w:val=\"24\"/></w:rPr>"),
        ("31", "toc 3", "", "<w:rPr><w:sz w:val=\"24\"/><w:szCs w:val=\"24\"/></w:rPr>"),
        ("a0", "Перечень нум", "", ""),
        ("a2", "Перечень марк",
         "<w:pPr><w:numPr><w:ilvl w:val=\"0\"/><w:numId w:val=\"2\"/></w:numPr></w:pPr>", ""),
        ("af3", "Название рисунка", "", ""),
        ("af4", "Название таблицы", "", ""),
        ("af5", "Текст в таблице", "", ""),
    )
    for sid, name, pPr, rPr in specs:
        if _has_style(doc, sid):
            continue
        _add_style(doc, para.format(ns=nsdecls("w"), sid=sid, name=name,
                                    pPr=pPr, rPr=rPr))
    if not _has_style(doc, "af0"):
        _add_style(doc, char.format(ns=nsdecls("w")))


def _install_defaults(doc) -> None:
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        return
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is not None:
        rpr = rpr_default.find(qn("w:rPr"))
        if rpr is None:
            rpr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
            rpr_default.append(rpr)
        for tag in ("w:sz", "w:szCs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = parse_xml(f'<{tag} {nsdecls("w")} w:val="24"/>')
                rpr.append(el)
            else:
                el.set(qn("w:val"), "24")
    ppr_default = defaults.find(qn("w:pPrDefault"))
    if ppr_default is not None:
        ppr = parse_xml(
            f'<w:pPr {nsdecls("w")}><w:spacing w:before="0" w:after="0" '
            f'w:line="360" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr>'
        )
        old = ppr_default.find(qn("w:pPr"))
        if old is not None:
            ppr_default.replace(old, ppr)
        else:
            ppr_default.append(ppr)


def _install_numbering(doc) -> None:
    try:
        numbering = doc.part.numbering_part._element
    except (AttributeError, NotImplementedError, KeyError):
        return
    for abstract in numbering.findall(qn("w:abstractNum")):
        for lvl in abstract.findall(qn("w:lvl")):
            # Привязка уровня к стилизованному ListBullet*/ListNumber*
            # мешает LibreOffice: маркеры/номера не отрисовываются,
            # когда нумерация задаётся явным numPr абзаца или стиля.
            pstyle = lvl.find(qn("w:pStyle"))
            if pstyle is not None:
                lvl.remove(pstyle)
        lvl0 = abstract.find(qn("w:lvl"))
        if lvl0 is None:
            continue
        fmt = lvl0.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet":
            continue
        lt = lvl0.find(qn("w:lvlText"))
        if lt is None:
            lt = parse_xml(f'<w:lvlText {nsdecls("w")} w:val="•"/>')
            lvl0.append(lt)
        else:
            lt.set(qn("w:val"), "•")
        rpr = lvl0.find(qn("w:rPr"))
        if rpr is not None:
            lvl0.remove(rpr)


def _build_front(doc) -> list:
    heading = doc.add_heading(_REV_HEADING, level=1)
    pPr = heading._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))

    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, _REV_HEADERS):
        cell.text = text
        cell.paragraphs[0].style = doc.styles["Текст в таблице"]
        for run in cell.paragraphs[0].runs:
            run.bold = True

    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}>'
        f"<w:sdtPr><w:id w:val=\"-7070707\"/>"
        f"<w:docPartObj><w:docPartGallery w:val=\"Table of Contents\"/>"
        f"<w:docPartUnique/></w:docPartObj></w:sdtPr>"
        f"<w:sdtContent>"
        f"<w:p><w:pPr><w:pStyle w:val=\"ae\"/></w:pPr>"
        f"<w:r><w:t>Содержание</w:t></w:r></w:p>"
        f"<w:p><w:pPr><w:pStyle w:val=\"11\"/>"
        f"<w:tabs>{_TOC_TAB}</w:tabs></w:pPr>"
        f"<w:r><w:fldChar w:fldCharType=\"begin\"/></w:r>"
        f"<w:r><w:instrText xml:space=\"preserve\"> TOC \\o \"1-4\" \\h \\z \\u </w:instrText></w:r>"
        f"<w:r><w:fldChar w:fldCharType=\"separate\"/></w:r>"
        f"</w:p>"
        f"<w:p><w:pPr><w:pStyle w:val=\"11\"/>"
        f"<w:tabs>{_TOC_TAB}</w:tabs></w:pPr>"
        f"<w:r><w:t>Пример</w:t></w:r>"
        f"<w:r><w:tab/></w:r>"
        f"<w:r><w:t>1</w:t></w:r>"
        f"</w:p>"
        f"<w:p><w:r><w:fldChar w:fldCharType=\"end\"/></w:r></w:p>"
        f"</w:sdtContent>"
        f"</w:sdt>"
    )
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    sect_pr.addprevious(sdt)

    return [heading._element, table._tbl, sdt]


def fresh_gost_doc() -> tuple:
    """Аналог ``generate_single._fresh_doc`` без primer.docx.

    Возвращает ``(doc, kept)`` — документ и три фронтовых элемента
    (заголовок и лист регистрации изменений, поле содержания).
    """
    doc = Document()
    _install_styles(doc)
    _install_defaults(doc)
    _install_numbering(doc)
    kept = _build_front(doc)
    return doc, kept


def template_warning_xml() -> str | None:
    """XML таблицы «Предупреждение» (вместо чтения из primer.docx)."""
    global _WARNING_XML
    if _WARNING_XML is None:
        borders = "".join(
            f'<w:{edge} w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        _WARNING_XML = (
            f'<w:tbl {nsdecls("w")}>'
            f"<w:tblPr><w:tblW w:type=\"dxa\" w:w=\"10205\"/>"
            f"<w:tblBorders>{borders}</w:tblBorders>"
            f"<w:tblLook w:val=\"04A0\" w:firstRow=\"1\" w:lastRow=\"0\" "
            f"w:firstColumn=\"1\" w:lastColumn=\"0\" w:noHBand=\"0\" w:noVBand=\"1\"/>"
            f"</w:tblPr>"
            f"<w:tblGrid><w:gridCol w:w=\"10205\"/></w:tblGrid>"
            f"<w:tr><w:tc><w:tcPr><w:tcW w:type=\"dxa\" w:w=\"10205\"/></w:tcPr>"
            f"<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Предупреждение</w:t></w:r></w:p>"
            f"</w:tc></w:tr>"
            f"<w:tr><w:tc><w:tcPr><w:tcW w:type=\"dxa\" w:w=\"10205\"/></w:tcPr>"
            f"<w:p/></w:tc></w:tr>"
            f"</w:tbl>"
        )
    return _WARNING_XML
