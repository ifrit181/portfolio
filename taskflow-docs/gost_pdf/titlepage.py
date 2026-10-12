"""Титульный лист по ГОСТ 19.104-78 (п. 3.1–3.2, прил. 3).

Поля 1–10 титульного листа (ГОСТ 19.104-78, п. 3.1); значения по умолчанию
пустые — см. `reference/matrix.md` § 4.1 и `reference/РЕКВИЗИТЫ.md`.
"""

from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


def _find_sect_break_para(body):
    for child in body:
        if child.tag != qn("w:p"):
            continue
        pPr = child.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            return child
    return None


def _para(doc, text="", *, size=12, bold=False, caps=False, align=None,
          before=0, after=0, line=1.5, tabs=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if tabs:
        for pos, kind in tabs:
            pf.tab_stops.add_tab_stop(pos, kind)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if caps:
            r.font.all_caps = True
    return p


def _nlines(text: str, width: int) -> int:
    text = " ".join(str(text or "").split())
    if not text:
        return 0
    return max(1, -(-len(text) // width))


def apply(doc, *, product_name: str, doc_type_title: str,
          oboznachenie: str = "", litera: str = "", otmetka_ucheta: str = "",
          stroka_izmenenij: str = "", lu_oboznachenie: str = "",
          ministerstvo: str = "", sheets: int | None = None,
          year: int | None = None, left30: bool = False):
    """Заменить блок до первого раздела документа на титульный лист."""
    body = doc.element.body
    sect_break = _find_sect_break_para(body)
    if sect_break is None:
        return
    for child in list(body):
        if child is sect_break:
            break
        body.remove(child)

    content_width = 170 if left30 else 180
    center_tab = Mm(content_width // 2)
    right_tab = Mm(content_width)

    line_pt = 12 * 1.5 * 1.17
    avail_pt = 257 * 2.8346

    def block_pt(text, size, width):
        return _nlines(text, width) * size * 1.5 * 1.17

    used_pt = 0.0
    if ministerstvo:
        used_pt += block_pt(ministerstvo, 12, 85)
    used_pt += 12 * 1.5 * 1.17
    if lu_oboznachenie:
        used_pt += block_pt(lu_oboznachenie, 12, 85)
    if oboznachenie:
        used_pt += block_pt(oboznachenie, 12, 85) + 6
    used_pt += block_pt(product_name, 16, 45)
    used_pt += block_pt(doc_type_title, 14, 55) + 6
    if sheets and sheets > 1:
        used_pt += 12 * 1.5 * 1.17
    if otmetka_ucheta:
        used_pt += block_pt(otmetka_ucheta, 11, 85)
    if stroka_izmenenij:
        used_pt += block_pt(stroka_izmenenij, 11, 85)
    used_pt += 12 * 1.5 * 1.17

    total_sp = int(max(4 * line_pt, avail_pt - used_pt - line_pt) // line_pt)
    s1 = max(1, round(total_sp * 0.34))
    s2 = max(1, round(total_sp * 0.24))
    s3 = max(1, round(total_sp * 0.20))
    s4 = max(1, total_sp - s1 - s2 - s3)
    s1 = max(1, total_sp - s2 - s3 - s4)

    paras = []
    left = WD_ALIGN_PARAGRAPH.LEFT
    center = WD_ALIGN_PARAGRAPH.CENTER
    right = WD_ALIGN_PARAGRAPH.RIGHT

    if ministerstvo:
        paras.append(_para(doc, ministerstvo, align=left))
    paras.append(_para(doc, "УТВЕРЖДЕНО", bold=True, align=left))
    if lu_oboznachenie:
        paras.append(_para(doc, lu_oboznachenie, align=left))
    if oboznachenie:
        paras.append(_para(doc, oboznachenie, align=left, before=6))
    paras.extend(_para(doc, align=left) for _ in range(s1))
    paras.append(_para(doc, product_name or "", size=16, bold=True, caps=True,
                       align=center))
    paras.append(_para(doc, doc_type_title, size=14, align=center, before=6))
    paras.extend(_para(doc, align=center) for _ in range(s2))
    if sheets and sheets > 1:
        paras.append(_para(doc, f"Листов {sheets}", align=right))
    paras.extend(_para(doc, align=right) for _ in range(s3))
    if otmetka_ucheta:
        paras.append(_para(doc, otmetka_ucheta, size=11, align=left))
    if stroka_izmenenij:
        paras.append(_para(doc, stroka_izmenenij, size=11, align=left))
    paras.extend(_para(doc, align=left) for _ in range(s4))

    bottom = _para(
        doc,
        tabs=[(center_tab, WD_TAB_ALIGNMENT.CENTER),
              (right_tab, WD_TAB_ALIGNMENT.RIGHT)],
    )
    bottom.add_run("\t" + (f"{year} г." if year else ""))
    if litera:
        bottom.add_run("\tЛитера " + litera)
    paras.append(bottom)

    for p in paras:
        sect_break.addprevious(p._element)
