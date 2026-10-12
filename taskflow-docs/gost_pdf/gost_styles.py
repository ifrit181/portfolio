"""Крестовые требования оформления (matrix.md § 2, § 3, § 7).

Реализация: структура документа (порядок частей, содержание, лист регистрации
изменений), шрифты и стили, поля листа, колонтитулы, подписи, перечни.
"""

from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt

FONT = "Liberation Sans"
MONO = "Liberation Mono"
IND_MM = 12.5

RENAME = {
    "Изменения в документе": "Лист регистрации изменений",
}

_REV_HEADING = "Изменения в документе"
_REV_HEADERS = {
    "Версия": "№ изм.",
    "Дата": "Дата",
    "Автор": "Подпись",
    "Описание": "Содержание изменения",
}


def _find_sect_break_para(body):
    for child in body:
        if child.tag != qn("w:p"):
            continue
        pPr = child.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            return child
    return None


def _find_toc_sdt(body):
    for el in body.iter(qn("w:sdt")):
        for it in el.iter(qn("w:instrText")):
            if it.text and it.text.strip().startswith("TOC"):
                return el
    return None


def _para_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def _set_para_text(p, new_text: str):
    first = True
    for t in p.iter(qn("w:t")):
        if first:
            t.text = new_text
            first = False
        else:
            t.text = ""


def _style(doc, sid: str):
    for s in doc.styles:
        if s.element.get(qn("w:styleId")) == sid:
            return s
    return None


def _set_pstyle(p, sid: str):
    pPr = p.find(qn("w:pPr"))
    if pPr is None:
        pPr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
        p.insert(0, pPr)
    ps = pPr.find(qn("w:pStyle"))
    if ps is None:
        ps = parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{sid}"/>')
        pPr.insert(0, ps)
    ps.set(qn("w:val"), sid)


def _style_cfg(st, *, size=None, bold=None, caps=None, align=None, line=None,
               before=None, after=None, ind_left=None, ind_first=None):
    if size is not None:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if caps is not None:
        st.font.all_caps = caps
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if line is not None:
        pf.line_spacing = line
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if ind_left is not None:
        pf.left_indent = Mm(ind_left)
    if ind_first is not None:
        pf.first_line_indent = Pt(ind_first)


def _structure_revision(doc, body, sect_break, left30: bool) -> bool:
    rev_h = None
    rev_tbl = None
    for p in body.findall(qn("w:p")):
        if _para_text(p).startswith(_REV_HEADING):
            rev_h = p
            break
    if rev_h is None:
        return False
    for t in doc.tables:
        if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip() == "Версия":
            rev_tbl = t
            break
    if rev_tbl is None:
        return False

    _set_para_text(rev_h, RENAME[_REV_HEADING])
    for cell in rev_tbl.rows[0].cells:
        old = cell.text.strip()
        new = _REV_HEADERS.get(old)
        if new:
            for p in cell._tc.findall(qn("w:p")):
                _set_para_text(p, new)

    content_mm = 170 if left30 else 180
    widths = [20, 25, 35, content_mm - 80]
    try:
        for col, mm_ in zip(rev_tbl.columns, widths):
            col.width = Mm(mm_)
        for row in rev_tbl.rows:
            for cell, mm_ in zip(row.cells, widths):
                cell.width = Mm(mm_)
        tbl = rev_tbl._tbl
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            total_tw = sum(Mm(w).twips for w in widths)
            tblW = tblPr.find(qn("w:tblW"))
            if tblW is None:
                tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:type="dxa" w:w="0"/>')
                anchor = tblPr.find(qn("w:tblStyle"))
                if anchor is not None:
                    anchor.addnext(tblW)
                else:
                    tblPr.insert(0, tblW)
            tblW.set(qn("w:w"), str(total_tw))
            tblW.set(qn("w:type"), "dxa")
        rev_tbl.autofit = False
    except (AttributeError, IndexError, ValueError):
        pass

    sectPr = body.find(qn("w:sectPr"))
    brk = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:br {nsdecls("w")} w:type="page"/></w:r></w:p>'
    )
    sectPr.addprevious(brk)
    sectPr.addprevious(rev_h)
    sectPr.addprevious(rev_tbl._tbl)
    return True


def _structure_annotation(doc, body, sect_break, annotation: str):
    sdt = _find_toc_sdt(body)
    break_p1 = None
    if sdt is not None:
        prev = sdt.getprevious()
        if (
            prev is not None
            and prev.tag == qn("w:p")
            and any(b.get(qn("w:type")) == "page" for b in prev.iter(qn("w:br")))
        ):
            break_p1 = prev

    if not annotation or not annotation.strip():
        if break_p1 is not None:
            break_p1.getparent().remove(break_p1)
        return

    paras = []
    head = doc.add_paragraph()
    pf = head.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(42)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = head.add_run("АННОТАЦИЯ")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.all_caps = True
    paras.append(head)

    for block in re.split(r"\n\s*\n", annotation.strip()):
        text = " ".join(ln.strip() for ln in block.splitlines() if ln.strip())
        if not text:
            continue
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(12)
        paras.append(p)

    anchor = break_p1 if break_p1 is not None else (sdt if sdt is not None else sect_break)
    for p in paras:
        anchor.addprevious(p._element)


def _structure_toc(doc, body, entries, left30: bool):
    sdt = _find_toc_sdt(body)
    if sdt is None:
        return
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return
    paras = content.findall(qn("w:p"))
    if not paras:
        return

    if _para_text(paras[0]) == "Содержание":
        _set_para_text(paras[0], "СОДЕРЖАНИЕ")

    for it in sdt.iter(qn("w:instrText")):
        if (it.text or "").strip().startswith("TOC"):
            it.text = ' TOC \\o "1-4" \\h \\z \\u '

    begin_p = None
    for p in paras:
        if any((it.text or "").strip().startswith("TOC") for it in p.iter(qn("w:instrText"))):
            begin_p = p
            break
    if begin_p is None:
        return
    end_p = None
    for p in reversed(paras):
        if p is begin_p:
            break
        if "end" in [f.get(qn("w:fldCharType")) for f in p.iter(qn("w:fldChar"))]:
            end_p = p
            break
    if end_p is None:
        return

    entry_paras = []
    started = False
    for p in paras:
        if p is begin_p:
            started = True
            continue
        if p is end_p:
            break
        if started:
            entry_paras.append(p)

    for p, (lvl, text) in zip(entry_paras, entries or []):
        _set_pstyle(p, "11" if lvl <= 2 else ("21" if lvl == 3 else "31"))
        new_text = RENAME.get(text)
        if new_text:
            ts = list(p.iter(qn("w:t")))
            if ts:
                ts[0].text = new_text
        if left30:
            for tab in p.iter(qn("w:tab")):
                if tab.get(qn("w:pos")):
                    tab.set(qn("w:pos"), str(Mm(165).twips))


def _fix_rfonts(el):
    for rf in el.iter(qn("w:rFonts")):
        ascii_v = (rf.get(qn("w:ascii")) or "").strip().lower()
        target = MONO if ascii_v in {"courier new", "consolas", "courier"} else FONT
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            key = qn("w:" + attr)
            if key in rf.attrib:
                del rf.attrib[key]
        rf.set(qn("w:ascii"), target)
        rf.set(qn("w:hAnsi"), target)
        rf.set(qn("w:cs"), target)


def _restyle(doc):
    center = WD_ALIGN_PARAGRAPH.CENTER
    left = WD_ALIGN_PARAGRAPH.LEFT
    justify = WD_ALIGN_PARAGRAPH.JUSTIFY

    for sid, size, bold, caps, align, indent in (
        ("1", 18, True, True, center, False),
        ("2", 16, True, True, center, False),
        ("3", 14, True, False, left, True),
        ("4", 14, False, False, left, True),
        ("5", 14, False, False, left, True),
        ("6", 14, False, False, left, True),
    ):
        st = _style(doc, sid)
        if st is None:
            continue
        st.font.name = FONT
        _style_cfg(
            st, size=size, bold=bold, caps=caps, align=align, line=1.5,
            before=0, after=size * 3,
            ind_left=IND_MM if indent else None,
            ind_first=0 if indent else None,
        )

    st = _style(doc, "a3")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=12, line=1.5, before=0, after=0, align=justify)

    for sid in ("a4", "a0", "a1", "a2", "af7"):
        st = _style(doc, sid)
        if st is None:
            continue
        st.font.name = FONT
        _style_cfg(st, size=12, line=1.5, before=0, after=0, align=justify)

    for sid, caps in (("11", False), ("21", False), ("31", False)):
        st = _style(doc, sid)
        if st is None:
            continue
        st.font.name = FONT
        _style_cfg(st, caps=caps, line=1.5, before=0, after=0)

    st = _style(doc, "ae")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=14, bold=True, caps=True, align=center, line=1.5,
                   before=0, after=42)

    st = _style(doc, "af3")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=10, align=center, line=1.5, before=6, after=6)

    st = _style(doc, "af4")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=10, align=left, line=1.5, before=6, after=6)

    st = _style(doc, "af5")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=10, align=justify, line=1.5, before=0, after=0)

    st = _style(doc, "af2")
    if st is not None:
        st.font.name = FONT
        _style_cfg(st, size=10, line=1.5)

    for sid in ("a8", "aa"):
        st = _style(doc, sid)
        if st is not None:
            st.font.name = FONT
            _style_cfg(st, size=10, before=0, after=0)

    body = doc.element.body
    target_tw = Mm(IND_MM).twips
    for p in body.iter(qn("w:p")):
        if any(a.tag == qn("w:sdt") for a in p.iterancestors()):
            continue
        if any(a.tag == qn("w:tbl") for a in p.iterancestors()):
            continue
        pPr = p.find(qn("w:pPr"))
        if pPr is not None:
            if pPr.find(qn("w:jc")) is not None:
                continue
            if pPr.find(qn("w:tabs")) is not None:
                continue
            if pPr.find(qn("w:numPr")) is not None:
                continue
            ps = pPr.find(qn("w:pStyle"))
            sid = ps.get(qn("w:val")) if ps is not None else None
            if sid not in (None, "", "a3", "a4"):
                continue
        if not _para_text(p):
            continue
        if pPr is None:
            pPr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
            p.insert(0, pPr)
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = parse_xml(f"<w:ind {nsdecls('w')}/>")
            anchor = None
            for name in ("w:rPr", "w:sectPr"):
                tail = pPr.find(qn(name))
                if tail is not None:
                    anchor = tail
                    break
            if anchor is not None:
                anchor.addprevious(ind)
            else:
                pPr.append(ind)
        ind.attrib.pop(qn("w:firstLineChars"), None)
        ind.set(qn("w:firstLine"), str(target_tw))

    cap_re = re.compile(r"^((?:Рисунок|Таблица)\s+[\d.]+)\s+-\s+")
    for p in body.iter(qn("w:p")):
        if any(a.tag == qn("w:sdt") for a in p.iterancestors()):
            continue
        for t in p.iter(qn("w:t")):
            if t.text:
                t.text = cap_re.sub(r"\1 — ", t.text)

    _patch_numbering(doc, target_tw)

    _fix_rfonts(body)
    _fix_rfonts(doc.styles.element)
    try:
        _fix_rfonts(doc.part.numbering_part._element)
    except (AttributeError, NotImplementedError, KeyError):
        pass
    for rel in doc.part.rels.values():
        if not (rel.reltype.endswith("/header") or rel.reltype.endswith("/footer")):
            continue
        part = rel.target_part
        el = getattr(part, "_element", None)
        if el is None:
            el = getattr(part, "element", None)
        if el is not None:
            _fix_rfonts(el)


def _patch_numbering(doc, target_tw: int):
    try:
        numbering = doc.part.numbering_part._element
    except (AttributeError, NotImplementedError, KeyError):
        return
    for abstract in numbering.findall(qn("w:abstractNum")):
        lvl0 = abstract.find(qn("w:lvl"))
        if lvl0 is None:
            continue
        pPr0 = lvl0.find(qn("w:pPr"))
        ind0 = pPr0.find(qn("w:ind")) if pPr0 is not None else None
        if ind0 is None:
            continue
        left0 = int(ind0.get(qn("w:left")) or 0)
        hang0 = int(ind0.get(qn("w:hanging")) or 0)
        delta = target_tw - (left0 - hang0)
        if delta == 0:
            continue
        for lvl in abstract.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            fmt = fmt.get(qn("w:val")) if fmt is not None else ""
            lt = lvl.find(qn("w:lvlText"))
            if (
                lt is not None
                and lt.get(qn("w:val"), "").endswith(".")
                and fmt in ("decimal", "lowerLetter", "upperLetter")
            ):
                lt.set(qn("w:val"), lt.get(qn("w:val"))[:-1] + ")")
            pPr = lvl.find(qn("w:pPr"))
            if pPr is None:
                continue
            ind = pPr.find(qn("w:ind"))
            if ind is not None:
                ind.set(qn("w:left"), str(int(ind.get(qn("w:left")) or 0) + delta))
            tabs = pPr.find(qn("w:tabs"))
            if tabs is not None:
                for tab in tabs.findall(qn("w:tab")):
                    pos = tab.get(qn("w:pos"))
                    if pos:
                        tab.set(qn("w:pos"), str(int(pos) + delta))


def _sections(doc, left30: bool):
    for s in doc.sections:
        s.page_width = Mm(210)
        s.page_height = Mm(297)
        s.left_margin = Mm(30 if left30 else 20)
        s.right_margin = Mm(10)
        s.top_margin = Mm(20)
        s.bottom_margin = Mm(20)
        s.header_distance = Mm(10)
        s.footer_distance = Mm(10)
        dg = s._sectPr.find(qn("w:docGrid"))
        if dg is not None:
            s._sectPr.remove(dg)


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _write_header(el, oboznachenie: str, center_pos: int):
    parts = [
        f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="a8"/>'
        f'<w:tabs><w:tab w:val="center" w:pos="{center_pos}"/></w:tabs></w:pPr>'
    ]
    if oboznachenie:
        parts.append(
            f'<w:r><w:t xml:space="preserve">{_xml_escape(oboznachenie)}</w:t></w:r>'
        )
    parts.append('<w:r><w:tab/></w:r>')
    parts.append(
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>1</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        "</w:p>"
    )
    new_p = parse_xml("".join(parts))
    for child in list(el):
        el.remove(child)
    el.append(new_p)


def _write_footer(el, text: str):
    new_p = parse_xml(
        f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="aa"/><w:jc w:val="right"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{_xml_escape(text)}</w:t></w:r></w:p>'
    )
    for child in list(el):
        el.remove(child)
    el.append(new_p)


def _clear_part(el):
    new_p = parse_xml(
        f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="aa"/></w:pPr></w:p>'
    )
    for child in list(el):
        el.remove(child)
    el.append(new_p)


def _kolontituly(doc, oboznachenie: str, stroka_izmenenij: str, left30: bool):
    hdr_ids = set()
    ftr_ids = set()
    sections = doc.sections
    for idx, section in enumerate(sections):
        sectPr = section._sectPr
        if idx == 0:
            for ref in list(sectPr.findall(qn("w:headerReference"))):
                sectPr.remove(ref)
            continue
        for ref in sectPr.findall(qn("w:headerReference")):
            if ref.get(qn("r:id")):
                hdr_ids.add(ref.get(qn("r:id")))
        for ref in sectPr.findall(qn("w:footerReference")):
            if ref.get(qn("r:id")):
                ftr_ids.add(ref.get(qn("r:id")))
    center_pos = Mm(85 if left30 else 90).twips
    for rel in doc.part.rels.values():
        is_hdr = rel.reltype.endswith("/header")
        is_ftr = rel.reltype.endswith("/footer")
        if not (is_hdr or is_ftr):
            continue
        part = rel.target_part
        el = getattr(part, "_element", None)
        if el is None:
            el = getattr(part, "element", None)
        if el is None:
            continue
        if is_hdr:
            if rel.rId in hdr_ids:
                _write_header(el, oboznachenie, center_pos)
            else:
                _clear_part(el)
        elif rel.rId in ftr_ids and stroka_izmenenij:
            _write_footer(el, stroka_izmenenij)
        else:
            _clear_part(el)


def apply(doc, *, entries, doc_type, oboznachenie: str = "",
          stroka_izmenenij: str = "", annotation: str = "", left30: bool = False):
    body = doc.element.body
    sect_break = _find_sect_break_para(body)

    _structure_revision(doc, body, sect_break, left30)
    _structure_annotation(doc, body, sect_break, annotation)
    _structure_toc(doc, body, entries, left30)
    _restyle(doc)
    _sections(doc, left30)
    _kolontituly(doc, oboznachenie, stroka_izmenenij, left30)

    doc.core_properties.subject = doc_type.title
